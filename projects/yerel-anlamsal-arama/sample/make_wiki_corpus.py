#!/usr/bin/env python3
"""Build a PUBLIC realistic Turkish corpus from Wikipedia to validate the finder
at scale (no private data). Writes .docx files + a queries.json of PARAPHRASED
queries (that avoid the title word, to test semantic recall, incl. confusable pairs).
"""
from __future__ import annotations

import datetime as dt
import json
import urllib.parse
import urllib.request
from pathlib import Path

import docx

HERE = Path(__file__).resolve().parent
OUT = HERE / "wiki"

# (title, filename, paraphrased query avoiding the title word)
ITEMS = [
    ("Kahve", "kahve", "sabah içilen kafeinli sıcak içecek"),
    ("Çay (içecek)", "cay", "demlenerek hazırlanan bitki yapraklı içecek"),
    ("Güneş", "gunes", "gök cismi, sistemimizin merkezindeki yıldız"),
    ("Ay", "ay", "dünyanın çevresinde dönen doğal uydu"),
    ("İstanbul", "istanbul", "iki kıtaya yayılan büyük Türk metropolü boğaz kenti"),
    ("Ankara", "ankara", "Türkiye'nin başkenti olan İç Anadolu şehri"),
    ("Fenerbahçe Spor Kulübü", "fenerbahce", "sarı lacivert renkli Kadıköy futbol kulübü"),
    ("Galatasaray Spor Kulübü", "galatasaray", "sarı kırmızı renkli İstanbul futbol kulübü"),
    ("Yapay zekâ", "yapay_zeka", "makinelerin insan gibi akıl yürütmesi teknolojisi"),
    ("Makine öğrenimi", "makine_ogrenimi", "verilerden örüntü öğrenen algoritmalar"),
    ("Osmanlı İmparatorluğu", "osmanli", "altı asır süren büyük Türk-İslam devleti"),
    ("Bizans İmparatorluğu", "bizans", "Doğu Roma'nın Konstantinopolis merkezli devamı"),
    ("Mimar Sinan", "mimar_sinan", "Süleymaniye'yi yapan ünlü Osmanlı yapı ustası"),
    ("Leonardo da Vinci", "davinci", "Mona Lisa'yı çizen Rönesans dâhisi"),
    ("Deprem", "deprem", "yer kabuğunun ani hareketiyle oluşan sarsıntı"),
    ("Yanardağ", "yanardag", "eriyik kaya ve lav püskürten dağ"),
    ("Fotosentez", "fotosentez", "bitkilerin ışıkla besin üretme süreci"),
    ("Bitcoin", "bitcoin", "merkeziyetsiz ilk dijital kripto para"),
    ("Blok zinciri", "blokzinciri", "dağıtık değiştirilemez kayıt defteri teknolojisi"),
    ("Nâzım Hikmet", "nazim", "memleket şiirleriyle tanınan Türk şair"),
    ("Orhan Pamuk", "orhan_pamuk", "Nobel kazanan Türk romancı"),
    ("Everest Dağı", "everest", "dünyanın en yüksek zirvesi Himalayalar'da"),
    ("Ağrı Dağı", "agri_dagi", "Türkiye'nin en yüksek dağı, Nuh efsanesiyle anılır"),
    ("Basketbol", "basketbol", "beşer kişilik iki takımın potaya top attığı oyun"),
    ("Antibiyotik", "antibiyotik", "bakteriyel enfeksiyonlara karşı kullanılan ilaç"),
    ("Aşı", "asi", "bağışıklık kazandıran koruyucu tıbbi uygulama"),
    ("Türk mutfağı", "turk_mutfagi", "kebap ve baklavayla ünlü yemek kültürü"),
    ("Küresel ısınma", "kuresel_isinma", "sera gazlarıyla gezegenin ısınması"),
    ("İnternet", "internet", "dünya çapında bağlı bilgisayar ağı"),
    ("Elektrik", "elektrik", "yüklü parçacıkların hareketiyle oluşan enerji"),
]


UA = "Atolye-finder/0.1 (+https://demle.me)"


def fetch_one(title: str) -> str:
    """Fetch a single article's plaintext; follow redirects, take the returned page."""
    q = urllib.parse.urlencode({
        "action": "query", "format": "json", "prop": "extracts",
        "explaintext": "1", "exlimit": "max", "redirects": "1", "titles": title})
    url = f"https://tr.wikipedia.org/w/api.php?{q}"
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=40) as r:
        data = json.loads(r.read())
    for p in data["query"]["pages"].values():
        if "extract" in p:
            return p["extract"]
    return ""


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    queries = []
    made = 0
    base = dt.datetime(2026, 1, 5, 9, 0)
    for idx, (title, fname, query) in enumerate(ITEMS):
        body = fetch_one(title)
        if not body:
            print(f"  ! bulunamadı: {title}")
            continue
        paras = [p.strip() for p in body.split("\n") if p.strip()][:8]
        body = "\n".join(paras)[:2500]
        d = docx.Document()
        d.add_heading(title, level=1)
        for p in body.split("\n"):
            d.add_paragraph(p)
        created = base + dt.timedelta(days=idx * 3)
        d.core_properties.created = created
        d.core_properties.modified = created
        d.core_properties.title = title
        d.core_properties.author = "Vikipedi"
        d.save(str(OUT / f"{fname}.docx"))
        queries.append({"query": query, "expect": fname})
        made += 1
    (HERE / "wiki_queries.json").write_text(
        json.dumps(queries, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\n{made} belge -> {OUT}/  ·  {len(queries)} sorgu -> wiki_queries.json")


if __name__ == "__main__":
    build()
