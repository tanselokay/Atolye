#!/usr/bin/env python3
# Generated with Claude Opus 4.8 (Anthropic). Review before use.
# Bu dosya Claude Opus 4.8 (Anthropic) ile üretilmiştir; kullanmadan önce gözden geçirin.
"""Generate a deliberately MESSY Turkish sample of real .docx / .pptx files.

Designed to stress the finder's two hard cases:
  1. Near-duplicate reports — same project, two versions differing by DATE + figures.
  2. Two same-name meeting notes — identical filename, different folders + dates,
     disambiguated only by "geçen haftaki" vs "bu haftaki".

Core.xml created/modified dates are set explicitly, anchored to 2026-08-22 (Sat):
  bu hafta   = Mon 2026-08-17 .. today
  geçen hafta= Mon 2026-08-10 .. Sun 2026-08-16
  dün        = 2026-08-21
  geçen ay   = Temmuz 2026
Files are .gitignored (may resemble real company docs).
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import docx
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PPt

HERE = Path(__file__).resolve().parent


def _set_dates(core, created: dt.datetime, modified: dt.datetime, author="Ayşe Yılmaz", title=""):
    core.created = created
    core.modified = modified
    core.author = author
    core.last_modified_by = author
    if title:
        core.title = title


def make_docx(relpath: str, heads_and_bodies, created, modified, author="Ayşe Yılmaz", title=""):
    d = docx.Document()
    for head, body in heads_and_bodies:
        if head:
            d.add_heading(head, level=1)
        for para in body:
            d.add_paragraph(para)
    _set_dates(d.core_properties, created, modified, author, title)
    out = HERE / relpath
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
    return out


def make_docx_table(relpath, heads_and_bodies, table_rows, created, modified, author="Ayşe Yılmaz", title=""):
    d = docx.Document()
    for head, body in heads_and_bodies:
        if head:
            d.add_heading(head, level=1)
        for para in body:
            d.add_paragraph(para)
    if table_rows:
        t = d.add_table(rows=0, cols=len(table_rows[0]))
        t.style = "Table Grid"
        for row in table_rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
    _set_dates(d.core_properties, created, modified, author, title)
    out = HERE / relpath
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
    return out


def make_pptx(relpath, slides, created, modified, author="Ayşe Yılmaz", title=""):
    prs = Presentation()
    blank = prs.slide_layouts[1]  # title + content
    for stitle, bullets, notes in slides:
        s = prs.slides.add_slide(blank)
        s.shapes.title.text = stitle
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    _set_dates(prs.core_properties, created, modified, author, title)
    out = HERE / relpath
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


D = dt.datetime  # shorthand


def build():
    made = []

    # --- HARD CASE 1: near-duplicate project reports (same project, 2 versions) ---
    made.append(make_docx_table(
        "raporlar/proje_durum_raporu.docx",
        [("Atlas Projesi Durum Raporu", [
            "Bu rapor Atlas müşteri portalı projesinin ilerlemesini özetler.",
            "Genel durum: proje planlanan takvimin bir miktar gerisinde ilerliyor.",
        ]),
         ("Özet", ["Tamamlanma oranı ve bütçe kullanımı aşağıdaki tabloda verilmiştir."])],
        [["Gösterge", "Değer"],
         ["Tamamlanma", "%60"],
         ["Harcanan bütçe", "1.2 milyon TL"],
         ["Açık riskler", "3"]],
        created=D(2026, 7, 15, 9, 0), modified=D(2026, 7, 15, 10, 30),
        title="Atlas Projesi Durum Raporu — Temmuz"))

    made.append(make_docx_table(
        "raporlar/proje_durum_raporu_guncel.docx",
        [("Atlas Projesi Durum Raporu", [
            "Bu rapor Atlas müşteri portalı projesinin ilerlemesini özetler.",
            "Genel durum: proje son iki haftada hız kazandı ve takvime yaklaştı.",
        ]),
         ("Özet", ["Tamamlanma oranı ve bütçe kullanımı aşağıdaki tabloda verilmiştir."])],
        [["Gösterge", "Değer"],
         ["Tamamlanma", "%85"],
         ["Harcanan bütçe", "1.5 milyon TL"],
         ["Açık riskler", "1"]],
        created=D(2026, 8, 20, 9, 0), modified=D(2026, 8, 20, 16, 0),
        title="Atlas Projesi Durum Raporu — Ağustos (güncel)"))

    # --- HARD CASE 2: two same-name meeting notes, different weeks + folders ---
    made.append(make_docx(
        "2026-08-12/toplanti_notlari.docx",
        [("Pazarlama Toplantısı Notları", [
            "Katılımcılar: Ayşe, Mehmet, Deniz.",
            "Üçüncü çeyrek kampanya bütçesi görüşüldü.",
            "Karar: sosyal medya harcaması %20 artırılacak.",
            "Aksiyon: Deniz yeni ajans tekliflerini toplayacak.",
        ])],
        created=D(2026, 8, 12, 14, 0), modified=D(2026, 8, 12, 15, 0),
        title="Pazarlama Toplantısı — 12 Ağustos"))

    made.append(make_docx(
        "2026-08-19/toplanti_notlari.docx",
        [("Pazarlama Toplantısı Notları", [
            "Katılımcılar: Ayşe, Mehmet, Can.",
            "Yeni ürün lansmanı zamanlaması konuşuldu.",
            "Karar: lansman Ekim başına çekildi.",
            "Aksiyon: Can basın bültenini hazırlayacak.",
        ])],
        created=D(2026, 8, 19, 14, 0), modified=D(2026, 8, 19, 15, 0),
        title="Pazarlama Toplantısı — 19 Ağustos"))

    # --- ordinary docx files (the everyday clutter) ---
    made.append(make_docx(
        "gizlilik_politikasi.docx",
        [("Kişisel Verilerin Korunması Politikası", [
            "Şirketimiz KVKK kapsamında kişisel verileri işler ve korur.",
            "Veriler yalnızca hizmetin gerektirdiği süre boyunca saklanır.",
            "Çalışanlar müşteri verilerini üçüncü taraflarla paylaşamaz.",
        ])],
        created=D(2026, 3, 2, 9, 0), modified=D(2026, 3, 2, 9, 0),
        title="Gizlilik Politikası"))

    made.append(make_docx(
        "izin_dilekcesi.docx",
        [("Yıllık İzin Dilekçesi", [
            "Sayın İnsan Kaynakları,",
            "1–5 Eylül 2026 tarihleri arasında yıllık iznimi kullanmak istiyorum.",
            "Gereğini bilgilerinize arz ederim.",
            "Mehmet Kaya",
        ])],
        created=D(2026, 8, 21, 11, 0), modified=D(2026, 8, 21, 11, 0),
        author="Mehmet Kaya", title="İzin Dilekçesi"))  # "dün" = 2026-08-21

    made.append(make_docx(
        "gecikmis_fatura_hatirlatma.docx",
        [("Ödeme Hatırlatması", [
            "Sayın Yetkili,",
            "15 Temmuz tarihli 2026-0417 numaralı faturanızın vadesi geçmiştir.",
            "Toplam 48.500 TL tutarındaki borcun en kısa sürede ödenmesini rica ederiz.",
        ])],
        created=D(2026, 8, 3, 10, 0), modified=D(2026, 8, 3, 10, 0),
        title="Gecikmiş Fatura Hatırlatması")) # letter

    made.append(make_docx(
        "tedarikci_sozlesme_teslimat.docx",
        [("Tedarik Sözleşmesi — Teslimat Şartları", [
            "Yüklenici, siparişleri sipariş tarihinden itibaren 30 gün içinde teslim eder.",
            "Geç teslimatta her gün için sözleşme bedelinin %0,5'i ceza uygulanır.",
            "Teslim yeri: İstanbul merkez depo.",
        ])],
        created=D(2026, 6, 10, 9, 0), modified=D(2026, 6, 12, 9, 0),
        title="Tedarik Sözleşmesi"))

    made.append(make_docx(
        "performans_degerlendirme.docx",
        [("Yıl Sonu Performans Değerlendirmesi", [
            "Çalışanın hedeflere ulaşma oranı ve gelişim alanları değerlendirilir.",
            "Bu dönem satış ekibi hedeflerin %112'sine ulaşmıştır.",
            "Gelişim önerisi: sunum becerileri için eğitim planlanmalı.",
        ])],
        created=D(2026, 1, 20, 9, 0), modified=D(2026, 1, 22, 9, 0),
        title="Performans Değerlendirme"))

    # --- a deck (pptx) ---
    made.append(make_pptx(
        "2027_butce_sunumu.pptx",
        [("2027 Bütçe Taslağı", ["Genel bakış", "Öncelikli yatırım alanları"],
          "Yönetim kuruluna sunulacak taslak."),
         ("Gelir Beklentisi", ["Toplam gelir hedefi: 42 milyon TL",
                               "Yeni pazar: %15 büyüme", "Mevcut müşteri: %8 büyüme"], ""),
         ("Yatırım Kalemleri", ["Ar-Ge: 6 milyon TL", "Pazarlama: 4 milyon TL",
                                "Altyapı ve bulut: 3 milyon TL"],
          "Ar-Ge payı geçen yıla göre iki kat.")],
        created=D(2026, 7, 28, 9, 0), modified=D(2026, 7, 29, 17, 0),
        title="2027 Bütçe Sunumu"))

    made.append(make_pptx(
        "urun_lansman_plani.pptx",
        [("Yeni Ürün Lansman Planı", ["Hedef kitle", "Zaman çizelgesi"],
          "Pazarlama ekibi için iç sunum."),
         ("Zaman Çizelgesi", ["Eylül: hazırlık", "Ekim: lansman", "Kasım: değerlendirme"], ""),
         ("Kanallar", ["Sosyal medya", "E-posta bülteni", "Etkinlik ve fuarlar"], "")],
        created=D(2026, 8, 10, 9, 0), modified=D(2026, 8, 18, 12, 0),
        title="Ürün Lansman Planı")) # bu hafta modified

    return made


if __name__ == "__main__":
    files = build()
    print(f"{len(files)} örnek dosya üretildi:")
    for f in files:
        print("  -", f.relative_to(HERE))
